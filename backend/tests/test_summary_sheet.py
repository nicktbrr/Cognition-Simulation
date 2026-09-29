"""Run with: cd backend && python -m unittest discover tests"""

import os
import tempfile
import unittest

import pandas as pd
from openpyxl import Workbook, load_workbook

from utils.summary_sheet import write_summary_sheet


GENDER = {
    'label': 'Gender',
    'category': 'demographics',
    'values': ['Man', 'Woman'],
    'distribution': {'type': 'categorical', 'percentages': {'Man': 60, 'Woman': 40}},
}

AGE = {
    'label': 'Age',
    'category': 'demographics',
    'values': ['18 - 82 years old'],
    'distribution': {'type': 'continuous', 'min': 18, 'max': 82, 'mean': 50, 'sd': 10.66},
}

# Saved before distributions existed: gets an even split.
LEGACY = {'label': 'Pet', 'category': 'lifestyle', 'values': ['Cat', 'Dog', 'None']}

ANXIETY = {
    'id': 'm-1',
    'title': 'Anxiety',
    'description': 'How anxious the persona feels.',
    'range': '1 - 5',
    'desiredValues': [
        {'value': 1, 'label': 'Not at all'},
        {'value': 5, 'label': 'Extremely'},
    ],
    'items': [{'id': 'i1', 'text': 'I feel tense'}, 'I worry a lot'],
    'citation': 'Spielberger (1983)',
}

STEPS = [
    {'id': 's1', 'label': 'Intro', 'previous': [], 'next': ['s2'],
     'instructions': 'Read the scenario.', 'temperature': 0.5, 'measures': [ANXIETY]},
    {'id': 's2', 'kind': 'measure', 'measure_id': 'm-1', 'label': 'Anxiety check',
     'previous': ['s1'], 'next': [], 'instructions': ANXIETY['description'],
     'measures': [ANXIETY]},
]

SPEC = {
    'title': 'Exam stress',
    'run_date': '2026-09-29 10:00',
    'model': 'gemini-2.0-flash',
    'num_personas': 20,
    'sample_name': 'Undergrads',
    'study_introduction': 'You are about to sit an exam.',
    'attributes': [GENDER, AGE, LEGACY],
    'steps': STEPS,
}


def rows_of(sheet):
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def row_starting(rows, first):
    return next(row for row in rows if row and row[0] == first)


class SummarySheetTest(unittest.TestCase):

    def setUp(self):
        workbook = Workbook()
        self.sheet = write_summary_sheet(workbook, SPEC)
        self.workbook = workbook
        self.rows = rows_of(self.sheet)
        self.cells = [value for row in self.rows for value in row if value is not None]

    def test_is_first_tab_with_every_section(self):
        self.assertEqual(self.workbook.sheetnames[0], 'Summary')
        for section in ['Simulation overview', 'Sample characteristics', 'Measures', 'Steps']:
            self.assertIn(section, self.cells)

    def test_overview(self):
        self.assertEqual(row_starting(self.rows, 'Model')[1], 'gemini-2.0-flash')
        self.assertEqual(row_starting(self.rows, 'Number of personas')[1], 20)
        self.assertEqual(row_starting(self.rows, 'Study introduction')[1], 'You are about to sit an exam.')

    def test_categorical_options_and_percentages(self):
        man = next(row for row in self.rows if row[3] == 'Man')
        woman = next(row for row in self.rows if row[3] == 'Woman')
        self.assertEqual(man[:5], ('Demographics', 'Gender', 'Categorical', 'Man', '60%'))
        # The characteristic is named once, on its first option.
        self.assertEqual(woman[:5], ('', '', '', 'Woman', '40%'))

    def test_continuous_range_mean_and_sd(self):
        age = next(row for row in self.rows if row[1] == 'Age')
        self.assertEqual(age[2:5], ('Continuous', '18 - 82 years old', 'mean 50, SD 10.7'))

    def test_legacy_attribute_gets_even_split(self):
        cat = next(row for row in self.rows if row[3] == 'Cat')
        self.assertEqual(cat[1], 'Pet')
        self.assertEqual(cat[4], '33.3%')

    def test_measure_anchors_items_and_citation(self):
        measure = row_starting(self.rows, 'Anxiety')
        self.assertEqual(measure[1], 'How anxious the persona feels.')
        self.assertEqual(measure[2], '1 = Not at all\n5 = Extremely')
        self.assertEqual(measure[3], '1. I feel tense\n2. I worry a lot')
        self.assertEqual(measure[4], 'Spielberger (1983)')

    def test_measure_used_on_two_steps_is_listed_once(self):
        self.assertEqual(sum(1 for row in self.rows if row[0] == 'Anxiety'), 1)

    def test_steps_by_label(self):
        intro = row_starting(self.rows, 'Intro')
        check = row_starting(self.rows, 'Anxiety check')
        self.assertEqual(intro[1:4], ('Step', '', 'Anxiety check'))
        self.assertEqual(check[1:4], ('Measure', 'Intro', ''))
        self.assertEqual(check[6], 'Anxiety')

    def test_empty_spec_still_renders(self):
        sheet = write_summary_sheet(Workbook(), {})
        cells = [v for row in rows_of(sheet) for v in row if v is not None]
        self.assertIn('No sample characteristics', cells)
        self.assertIn('No measures', cells)

    def test_first_tab_alongside_pandas_sheets(self):
        # dataframe_to_excel writes the Summary through pandas' own workbook.
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'out.xlsx')
            with pd.ExcelWriter(path, engine='openpyxl') as writer:
                write_summary_sheet(writer.book, SPEC)
                pd.DataFrame({'ID': [1]}).to_excel(writer, sheet_name='Personas', index=False)
            self.assertEqual(load_workbook(path).sheetnames, ['Summary', 'Personas'])


if __name__ == '__main__':
    unittest.main()
